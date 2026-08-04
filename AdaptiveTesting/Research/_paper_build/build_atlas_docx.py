"""Build a minimal, graph-forward .docx covering ONLY the MCQ ATLAS Recreation experiment.

Run: uv run --with python-docx python build_atlas_docx.py
Output: ATLAS_Recreation_MCQ.docx (does not touch the full-paper docx).
"""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
MCQ = os.path.abspath(os.path.join(HERE, "..", "01_MCQ_ATLAS"))
FIG = os.path.join(MCQ, "figures")
SELF = os.path.join(MCQ, "data", "atlas_selfresp", "figures")
OUT = os.path.join(HERE, "ATLAS_Recreation_MCQ.docx")
IMG_W = Inches(5.5)

embedded, missing = [], []
doc = Document()

# tighten default spacing for a plain look
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)


def para(text, size=11, bold=False, italic=False, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.bold, r.italic, r.font.size = bold, italic, Pt(size)
    return p


def heading(text):
    para(text, size=13, bold=True, space_after=4)


def figure(path, caption):
    if not os.path.exists(path):
        missing.append(os.path.basename(path))
        para("[missing figure: %s]" % os.path.basename(path), size=9, italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=IMG_W)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run(caption)
    r.italic, r.font.size = True, Pt(9)
    r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    embedded.append(os.path.basename(path))


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i].paragraphs[0]
        rr = c.add_run(h)
        rr.bold, rr.font.size = True, Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cp = cells[i].paragraphs[0]
            rr = cp.add_run(str(v))
            rr.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


# ---- Title + intro ----------------------------------------------------------
title = doc.add_paragraph()
tr = title.add_run("ATLAS Recreation (MCQ): item-bank transfer, prompt-matched scoring, and OpenLM replication")
tr.bold, tr.font.size = True, Pt(16)
title.paragraph_format.space_after = Pt(8)

para(
    "We rebuild the ATLAS adaptive-testing (CAT) pipeline on our own data. Every correlation is the "
    "Pearson r between the CAT-predicted benchmark score and the true full-benchmark accuracy on held-out "
    "models. Data sources: (1) ATLAS's published 3PL item bank, (2) ATLAS's own ARC test responses "
    "(prompt-matched), and (3) OpenLM per-question model responses for the multi-benchmark replication."
)

# ---- Part 1: transfer -------------------------------------------------------
heading("1. A published bank transfers on ARC; a range-restricted refit does not")
para(
    "ATLAS-style CAT on ARC, scored on the same 60 held-out models, our 0-shot loglikelihood responses. "
    "The published wide-pool ATLAS 3PL bank recovers held-out ARC accuracy at r=0.83 (SE\u22640.2) using "
    "about 21 of 1,172 items, and r=0.74 at SE\u22640.3 (~10 items). Refitting the identical pipeline on only "
    "the 0.5\u20137B model slice (~1,686 models, ~95% at 7B) drops transfer to r=0.59, even though MAE is "
    "almost unchanged (~0.147). The loss is in ranking correlation, not average error: pool parameter-range "
    "coverage matters more than matching the target range."
)
table(
    ["Calibration pool", "SE stop", "Pearson r", "MAE", "Mean items"],
    [
        ["Published ATLAS bank", "0.2", "0.830", "0.147", "21.1"],
        ["Published ATLAS bank", "0.3", "0.738", "0.172", "9.9"],
        ["Refit, 0.5\u20137B", "0.2", "0.589", "0.147", "12.5"],
        ["Refit, 0.5\u20137B", "0.3", "0.585", "0.150", "8.3"],
    ],
)
figure(os.path.join(FIG, "atlas_arc_heldout_se0.2.png"),
       "Published ATLAS bank, SE\u22640.2: p-IRT predicted vs actual held-out ARC accuracy. r=0.83, MAE 0.147, ~21 items.")
figure(os.path.join(FIG, "atlas_arc_heldout_se0.3.png"),
       "Published ATLAS bank, SE\u22640.3: r=0.74, MAE 0.172, ~10 items.")
figure(os.path.join(FIG, "atlas_arc_0p5_7b_heldout_se0.2.png"),
       "Refit on 0.5\u20137B slice, SE\u22640.2: r=0.59, MAE 0.147 \u2014 same average error, worse ranking.")
figure(os.path.join(FIG, "atlas_arc_0p5_7b_heldout_se0.3.png"),
       "Refit on 0.5\u20137B slice, SE\u22640.3: r=0.59, MAE 0.150.")

# ---- Part 2: self-response --------------------------------------------------
heading("2. Prompt-matched responses recover near-ceiling transfer")
para(
    "The transfer test above has a scoring confound: our 0-shot responses were scored against a bank ATLAS "
    "calibrated on 25-shot leaderboard responses. Replaying the same 3PL CAT on ATLAS's own ARC test response "
    "matrix (aligned by item index; 650 bank items align) removes the mismatch. On a seed-7 sample of 60 models "
    "from ATLAS's 417-model ARC test set, the published bank reaches r=0.91 in under 10 items and MAE drops "
    "about 5x, from 0.147 to 0.031 at SE\u22640.2. Most of the earlier gap was a scoring/prompt mismatch, not an "
    "IRT transfer failure."
)
table(
    ["Held-out responses", "SE stop", "Pearson r", "MAE", "Mean items"],
    [
        ["Our 0-shot responses", "0.2", "0.830", "0.147", "21.1"],
        ["Our 0-shot responses", "0.3", "0.738", "0.172", "9.9"],
        ["ATLAS own responses", "0.2", "0.915", "0.031", "13.0"],
        ["ATLAS own responses", "0.3", "0.906", "0.034", "9.2"],
    ],
)
figure(os.path.join(SELF, "atlas_selfresp_arc_se0.2.png"),
       "ATLAS's own (prompt-matched) ARC responses, SE\u22640.2: r=0.91, MAE 0.031 (~5x lower than the 0-shot replay).")
figure(os.path.join(SELF, "atlas_selfresp_arc_se0.3.png"),
       "ATLAS's own responses, SE\u22640.3: r=0.91, MAE 0.034, ~9 items.")

# ---- Part 3: OpenLM replication ---------------------------------------------
heading("3. Replicating the ATLAS error analyses on OpenLM Leaderboard v2 tasks")
para(
    "Using OpenLM per-question model responses (0.2\u20137B models) we run the frozen 3PL banks + Fisher-information "
    "p-IRT selector on five tasks ATLAS never covered, at ATLAS's own SE thresholds (0.1/0.2/0.3), single "
    "calibration on a fixed seed-7 90/10 split (~110 held-out; math 90). On IFEval, MATH, and GPQA the accuracy "
    "MAE (0.016\u20130.070) matches ATLAS's 0.02\u20130.05 band and falls as SE tightens; adaptive selection reaches a "
    "given SE in 1.6x\u20135.6x fewer items than random."
)
para(
    "One preprocessing step decides these numbers: dropping non-positive-discrimination items (a\u22640, mirt "
    "failures / reverse-scored). The GPQA, MuSR, and BBH banks are 51% / 43% / 31% such items. Keeping them "
    "collapses GPQA to r\u22480; dropping them recovers r=0.74 (SE\u22640.3). Under the same filter MuSR recovers to "
    "~0.75 (not the clean negative control it first appeared). BBH stays weakest (r\u22480.68): rank-correlated but "
    "biased, because a single-factor 3PL underfits its multi-subtask mixture."
)
table(
    ["Benchmark", "r (0.1/0.2/0.3)", "acc-MAE (0.1/0.2/0.3)", "theta-MAE (0.1/0.2/0.3)", "items @0.3 adpt/rand"],
    [
        ["IFEval", "0.955/0.927/0.921", "0.034/0.047/0.050", "0.070/0.160/0.207", "18 / 62"],
        ["GPQA", "0.710/0.757/0.737", "0.034/0.066/0.070", "0.104/0.357/0.407", "13 / 46"],
        ["MATH", "0.948/0.892/0.878", "0.016/0.023/0.024", "0.097/0.168/0.208", "77 / 148"],
        ["MuSR", "0.893/0.770/0.746", "0.075/0.091/0.095", "0.362/0.629/0.707", "14 / 76"],
        ["BBH", "0.695/0.677/0.674", "0.109/0.112/0.114", "0.711/0.719/0.684", "8 / 13"],
    ],
)
para("Cross-benchmark error curves:", size=11, bold=True, space_after=4)
figure(os.path.join(FIG, "atlasrep_mae_vs_se.png"),
       "Accuracy MAE vs SE stopping threshold across the five OpenLM tasks; MAE falls as SE tightens.")
figure(os.path.join(FIG, "atlasrep_items_vs_se.png"),
       "Mean items vs SE: Fisher-info adaptive selection vs a random-item baseline (1.6x\u20135.6x fewer items).")
figure(os.path.join(FIG, "atlasrep_theta_mae_vs_se.png"),
       "Ability (theta) MAE vs SE across the five benchmarks.")
para("Per-benchmark p-IRT predicted vs actual accuracy (SE\u22640.2):", size=11, bold=True, space_after=4)
figure(os.path.join(FIG, "atlasrep_ifeval_pirt_vs_actual_se_0.2.png"),
       "IFEval, SE\u22640.2: r=0.93, acc-MAE 0.047 \u2014 tracks the y=x line.")
figure(os.path.join(FIG, "atlasrep_gpqa_pirt_vs_actual_se_0.2.png"),
       "GPQA, SE\u22640.2: r=0.76 after dropping a\u22640 items (r\u22480 if kept).")
figure(os.path.join(FIG, "atlasrep_math_pirt_vs_actual_se_0.2.png"),
       "MATH, SE\u22640.2: r=0.89, acc-MAE 0.023 \u2014 tracks the y=x line.")
figure(os.path.join(FIG, "atlasrep_musr_pirt_vs_actual_se_0.2.png"),
       "MuSR, SE\u22640.2: r=0.77 after the a>0 filter \u2014 not a clean negative control.")
figure(os.path.join(FIG, "atlasrep_bbh_pirt_vs_actual_se_0.2.png"),
       "BBH, SE\u22640.2: r=0.68 \u2014 rank-correlated but biased (unidimensional 3PL underfits a multi-subtask bank).")

doc.save(OUT)
print("SAVED:", OUT)
print("SIZE_BYTES:", os.path.getsize(OUT))
print("EMBEDDED (%d):" % len(embedded), embedded)
print("MISSING (%d):" % len(missing), missing)
